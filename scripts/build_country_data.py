"""Build a public-only, versioned country dataset from World Bank workbooks.
Requires openpyxl for release preparation only; the agents have no dependency.
"""
import argparse,hashlib,json,math,zipfile
from pathlib import Path
import openpyxl
WGI_URL='https://datacatalogfiles.worldbank.org/ddh-published/0038026/DR0095947/WGI%202026%20Governance%20Estimates%20and%20Scores%20%281996-2025%29.xlsx'
BREADY_URL='https://www.worldbank.org/content/dam/sites/b-ready/documents/excel/B-READY_ALL_DATA_2025.zip'
WGI_LABELS={'va':'公众参与与问责','pv':'政治稳定','ge':'政府效能','rq':'监管质量','rl':'法治','cc':'腐败控制'}
PILLARS=['规制框架','公共服务','运营效率']
TOPICS=['企业开办','经营场所','公用事业','劳工','金融服务','国际贸易','税收','争端解决','市场竞争','企业破产']
def build(wgi_file,bready_file,as_of):
    countries={}
    w=openpyxl.load_workbook(wgi_file,read_only=True,data_only=True)
    for dim in WGI_LABELS:
        for n,r in enumerate(w[dim].values,1):
            if r[5]!=2025: continue
            code=r[2]; c=countries.setdefault(code,{'code':code,'name':r[1],'region':r[3],'income_group':r[4],'wgi':{},'bready':None})
            score=r[12]
            if isinstance(score,(int,float)) and math.isfinite(score):
                c['wgi'][dim]={'label':WGI_LABELS[dim],'value':score,'unit':'score_0_100','period':'2025','ci90':[r[14],r[15]],'source_count':r[7],'locator':f'{dim}!M{n} (score); O{n}:P{n} (90% CI)'}
    w=openpyxl.load_workbook(bready_file,read_only=True,data_only=True)
    for n,r in enumerate(w['00_B-READY_Pillar_Score'].values,1):
        if r[1] not in countries or not all(isinstance(x,(int,float)) for x in r[2:5]):continue
        countries[r[1]]['bready']={'period':'2025','pillars':[{'label':label,'value':val,'unit':'score_0_100','locator':f'00_B-READY_Pillar_Score!{col}{n}'} for label,val,col in zip(PILLARS,r[2:5],'CDE')],'topics':[]}
    for sheet,label in zip(w.sheetnames[2:],TOPICS):
        for n,r in enumerate(w[sheet].values,1):
            if r[1] in countries and countries[r[1]]['bready'] is not None and isinstance(r[2],(int,float)):
                countries[r[1]]['bready']['topics'].append({'label':label,'value':r[2],'unit':'score_0_100','locator':f'{sheet}!C{n}'})
    return {'schema':'country-baseline/v1','method_version':'public-reference-2026-v1','as_of':as_of,'available_from':as_of+'T00:00:00+00:00','status':'research_reference','score_semantics':'高分表示所选治理与营商指标表现较好；不能推出具体企业盈利或安全程度。',
    'formula':{'combined':'0.60 * mean(B-READY three pillars) + 0.40 * mean(WGI six scores)','wgi_only':'mean(WGI six scores); observation only, not comparable to combined score','missing':'No imputation; fewer than six WGI dimensions => no reference total','events':'No automatic event adjustment or simulated daily changes'},
    'sources':[{'id':'wgi2026','name':'World Bank Worldwide Governance Indicators, 2026 Update','period':'2025','url':'https://datacatalog.worldbank.org/infrastructure-data/search/dataset/0038026/worldwide-governance-indicators','download_url':WGI_URL,'sha256':hashlib.sha256(Path(wgi_file).read_bytes()).hexdigest(),'license':'CC BY 4.0 + World Bank dataset terms','license_url':'https://www.worldbank.org/en/about/legal/terms-of-use-for-datasets'},
    {'id':'bready2025','name':'World Bank Business Ready 2025','period':'2025','url':'https://www.worldbank.org/en/businessready/data','download_url':BREADY_URL,'sha256':hashlib.sha256(Path(bready_file).read_bytes()).hexdigest(),'hash_scope':'01_B-READY-2025-PILLAR-TOPIC-SCORES.xlsx inside ZIP','license':'CC BY 4.0 + World Bank dataset terms','license_url':'https://data.worldbank.org/indicator/IC.BRE.P1.RF'}],
    'limitations':['本项目公开的研究参考合成，不是世界银行官方综合指数；不发布国家排名。', 'WGI反映感知调查与专家评估；源级置信区间不等于合成总分置信区间。', '两套来源概念不同，合成权重属于研究选择，未作预测效度校准。', '年度结构指标与近期事件分开；细微分差不能推出统计显著或企业经营结果。'],
    'countries':list(sorted(countries.values(),key=lambda x:x['name']))}
if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--wgi',required=True);p.add_argument('--bready',required=True);p.add_argument('--as-of',required=True);p.add_argument('--output',required=True);a=p.parse_args()
    data=build(a.wgi,a.bready,a.as_of);Path(a.output).write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n',encoding='utf-8');print(json.dumps({'countries':len(data['countries']),'bready':sum(c['bready'] is not None for c in data['countries'])}))
